from flask import Flask, render_template, jsonify, request, send_file
import json
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from database import Database

app = Flask(__name__)
db = Database()

WAVE_LABELS = {
    "senoidal": "Senoidal",
    "cuadrada": "Cuadrada",
    "triangular": "Triangular",
    "sierra": "Sierra",
    "ruido": "Ruido",
}

SEA_LABELS = {
    "ola suave": "Ola Suave",
    "ola rompiente": "Ola Rompiente",
    "ola de rebote": "Ola de Rebote",
    "ola de tormenta": "Ola de Tormenta",
    "mar agitado": "Mar Agitado",
}

WAVE_COLORS = {
    "senoidal": "#36a2eb",
    "cuadrada": "#ff6384",
    "triangular": "#ffce56",
    "sierra": "#4bc0c0",
    "ruido": "#9966ff",
}

SEA_COLORS = {
    "ola suave": "#36a2eb",
    "ola rompiente": "#ff6384",
    "ola de rebote": "#ffce56",
    "ola de tormenta": "#4bc0c0",
    "mar agitado": "#9966ff",
}


@app.route("/")
def index():
    records = db.get_recent(100)
    stats = _compute_stats(records)
    return render_template("index.html", stats=stats, records=records)


@app.route("/api/readings")
def api_readings():
    since_id = request.args.get("since_id", type=int)
    if since_id is not None:
        records = db.get_since_id(since_id)
    else:
        records = db.get_recent(500)
    data = []
    for r in records:
        data.append({
            "id": r["id"],
            "timestamp": r["timestamp"],
            "wave_type": r["wave_type"],
            "sea_type": r["sea_type"],
            "channel": r["channel"],
            "amplitude": r["amplitude"],
            "frequency": r["frequency"],
            "period": r["period"],
        })
    return jsonify(data)


@app.route("/api/latest")
def api_latest():
    records = db.get_recent(1)
    if not records:
        return jsonify(None)
    r = records[0]
    waveform_data = None
    if r["data"] is not None:
        try:
            waveform_data = json.loads(r["data"])
        except (json.JSONDecodeError, TypeError):
            waveform_data = None
    return jsonify({
        "id": r["id"],
        "timestamp": r["timestamp"],
        "wave_type": r["wave_type"],
        "sea_type": r["sea_type"],
        "channel": r["channel"],
        "amplitude": r["amplitude"],
        "frequency": r["frequency"],
        "period": r["period"],
        "data": waveform_data,
    })


@app.route("/api/export")
def api_export():
    records = db.get_all()
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="36A2EB", end_color="36A2EB", fill_type="solid")

    # --- Hoja 1: Resumen ---
    ws1 = wb.active
    ws1.title = "Resumen"
    headers1 = ["ID", "Fecha/Hora", "Canal", "Onda", "Mar",
                "Amplitud (V)", "Frecuencia (Hz)", "Periodo (s)"]
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, r in enumerate(records, 2):
        ws1.cell(row=row_idx, column=1, value=r["id"])
        ws1.cell(row=row_idx, column=2, value=r["timestamp"])
        ws1.cell(row=row_idx, column=3, value=r["channel"])
        ws1.cell(row=row_idx, column=4, value=WAVE_LABELS.get(r["wave_type"], r["wave_type"]))
        ws1.cell(row=row_idx, column=5, value=SEA_LABELS.get(r["sea_type"], r["sea_type"]))
        if r["amplitude"] is not None:
            ws1.cell(row=row_idx, column=6, value=round(r["amplitude"], 4))
        if r["frequency"] is not None:
            ws1.cell(row=row_idx, column=7, value=round(r["frequency"], 2))
        if r["period"] is not None:
            ws1.cell(row=row_idx, column=8, value=round(r["period"], 6))

    for col_idx in range(1, 9):
        ws1.column_dimensions[get_column_letter(col_idx)].width = 18

    # --- Hoja 2: Muestras ---
    ws2 = wb.create_sheet("Muestras")
    rows_with_data = []
    max_samples = 0
    for r in records:
        if r["data"] is None:
            continue
        try:
            samples = json.loads(r["data"])
        except (json.JSONDecodeError, TypeError):
            continue
        if not samples:
            continue
        rows_with_data.append((r["id"], r["period"], samples))
        max_samples = max(max_samples, len(samples))

    headers2 = ["ID"] + [f"t_{i} (s)" for i in range(max_samples)]
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, (w_id, period, samples) in enumerate(rows_with_data, 2):
        ws2.cell(row=row_idx, column=1, value=w_id)
        n = len(samples)
        if period and n > 1:
            dt = period / (n - 1)
        else:
            dt = 0
        for i, val in enumerate(samples):
            cell = ws2.cell(row=row_idx, column=2 + i, value=round(val, 4))
        for i in range(max_samples):
            ts_cell = ws2.cell(row=1, column=2 + i)
            if dt:
                ts_cell.value = f"t={round(i * dt, 6)}s"

    for col_idx in range(1, max_samples + 2):
        ws2.column_dimensions[get_column_letter(col_idx)].width = 12

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="waveforms.xlsx",
    )


@app.route("/api/clear", methods=["POST"])
def api_clear():
    db.clear()
    return jsonify({"status": "ok"})


@app.route("/api/stats")
def api_stats():
    records = db.get_recent(500)
    return jsonify(_compute_stats(records))


def _compute_stats(records):
    total = len(records)
    if total == 0:
        return {
            "total": 0,
            "last_reading": None,
            "most_common_wave": None,
            "most_common_sea": None,
            "wave_distribution": {},
            "sea_distribution": {},
            "avg_frequency": 0,
            "avg_amplitude": 0,
            "channel_counts": {},
        }

    wave_counts = {}
    sea_counts = {}
    channel_counts = {}
    freqs = []
    amps = []

    for r in records:
        wt = r["wave_type"]
        st = r["sea_type"]
        ch = r["channel"]
        wave_counts[wt] = wave_counts.get(wt, 0) + 1
        sea_counts[st] = sea_counts.get(st, 0) + 1
        channel_counts[ch] = channel_counts.get(ch, 0) + 1
        if r["frequency"]:
            freqs.append(r["frequency"])
        if r["amplitude"]:
            amps.append(r["amplitude"])

    most_common_wave = max(wave_counts, key=wave_counts.get) if wave_counts else None
    most_common_sea = max(sea_counts, key=sea_counts.get) if sea_counts else None

    return {
        "total": total,
        "last_reading": records[0]["timestamp"] if records else None,
        "most_common_wave": WAVE_LABELS.get(most_common_wave, most_common_wave),
        "most_common_sea": SEA_LABELS.get(most_common_sea, most_common_sea),
        "wave_distribution": wave_counts,
        "sea_distribution": sea_counts,
        "avg_frequency": sum(freqs) / len(freqs) if freqs else 0,
        "avg_amplitude": sum(amps) / len(amps) if amps else 0,
        "channel_counts": channel_counts,
    }


if __name__ == "__main__":
    app.run(debug=True, port=5000)